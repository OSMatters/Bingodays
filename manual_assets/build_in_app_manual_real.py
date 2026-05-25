from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "manual_assets"
OUT = ROOT / "Bingodays_产品内使用手册_第一人称版.docx"


def set_run(run, size=10.5, bold=False, color=None):
    run.font.name = "PingFang SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def configure(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)

    for name, size in [("Heading 1", 15), ("Heading 2", 12.5)]:
        style = doc.styles[name]
        style.font.name = "PingFang SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("2E74B5")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)


def para(doc, text: str, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run(r, size=size)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run(r)


def title(doc):
    p = doc.add_paragraph()
    r = p.add_run("Bingodays 使用手册")
    set_run(r, 24, True, "17202C")

    p = doc.add_paragraph()
    r = p.add_run("按照以下步骤，可以创建属于自己的 Bingo 面板。")
    set_run(r, 11, False, "555555")


def image(doc, filename: str, caption: str, width=2.62):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    set_run(r, 8.8, False, "666666")
    r.italic = True


def marker_table(doc, rows):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    tbl.autofit = False

    for mark, text in rows:
        cells = tbl.add_row().cells
        cells[0].width = Inches(0.62)
        cells[1].width = Inches(5.48)
        shade(cells[0], "E8EEF5")

        cells[0].text = ""
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(mark))
        set_run(r0, 9.5, True, "2E74B5")

        cells[1].text = ""
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(text)
        set_run(r1, 9.4)

        cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cells[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    doc.add_paragraph()


def build():
    doc = Document()
    configure(doc)
    title(doc)

    doc.add_heading("1. 我先长按格子，进入任务编辑页", level=1)
    para(doc, "我可以从首页长按任意任务格，打开单个格子的编辑页。这里用来设置这一格要做什么、什么时候做、是否需要倒计时，以及是否要重复出现。")
    image(doc, "real_task_edit_marked.png", "任务编辑页：长按格子后，设置这一格的任务")
    marker_table(doc, [
        ("1", "任务名称：我在这里输入要完成的具体动作，也可以点击麦克风用语音输入。建议写成小而明确的任务，例如「喝一杯水」「阅读 10 页」，不要写得太笼统。"),
        ("2", "开始时间与结束时间：如果我希望任务只在某个时间段出现，可以设置开始和结束时间。到了开始时间后任务进入可执行状态；如果过了结束时间还没完成，会进入超时处理。"),
        ("3", "强制完成：如果这件事非常重要，我可以开启强制完成。开启后，系统会优先要求我先完成这个任务，再继续处理其他格子。"),
        ("4", "重复日期：我可以设置这个任务每天出现、只在指定星期出现，或只作为当天任务使用。适合晨间流程、运动习惯、学习计划等重复任务。"),
        ("5", "任务倒计时：我可以给任务设置预计完成时间，例如 15 分钟、30 分钟或自定义分钟数。倒计时结束后，系统会提示我选择已完成、延期，或删除这个任务。"),
        ("6", "保存：确认任务名称、时间、重复和倒计时设置后，点击保存。保存后回到首页，这一格就会显示新的任务。"),
    ])

    doc.add_heading("2. 我回到面板，点击格子表示完成", level=1)
    para(doc, "任务创建好后，我回到 Bingo 面板。每完成一件事，就点击对应的格子，格子会变成完成状态。只要完成的格子连成一行、一列或一条对角线，就会形成 Bingo。")
    image(doc, "real_home_marked.png", "首页面板：点击格子完成任务，连成线就是 Bingo")
    marker_table(doc, [
        ("1", "菜单入口：我可以进入设置、日记、小工具说明等页面。"),
        ("2", "积分入口：我可以查看目前积分，并进入贴纸和自定义奖励页面。"),
        ("3", "当前面板：这里显示我正在使用的面板名称，点击后可以切换或管理面板。"),
        ("4", "添加面板：如果我想为学习、家务、睡前流程建立不同面板，可以从这里添加。"),
        ("5", "任务格子：点击表示完成；再次点击可以取消完成；长按可以重新编辑这个格子。"),
        ("6", "编辑入口：进入面板编辑相关功能，例如快速编辑、调整格子大小、随机排列和清空面板。"),
    ])

    doc.add_heading("3. 我用 Bingo 积分兑换贴纸或自定义奖励", level=1)
    para(doc, "完成任务和达成 Bingo 后，我会获得积分。积分可以用来兑换贴纸，也可以用来兑换我自己创建的奖励卡片，让完成任务变得更有反馈。")
    image(doc, "real_points_marked.png", "积分页：用积分兑换贴纸，也可以创建自己的奖励卡片")
    marker_table(doc, [
        ("1", "我的积分：这里显示当前可使用的积分。兑换奖励时会消耗积分。"),
        ("2", "贴纸兑换：我可以用积分兑换贴纸，兑换后的贴纸可以添加到首页作为装饰。"),
        ("3", "兑换历史：我可以查看以前兑换过什么，方便回顾积分使用记录。"),
        ("4", "贴纸卡片：点击贴纸卡片可以查看兑换状态；积分足够时可以直接兑换。"),
        ("5", "自定义奖励卡片：我可以创建自己的奖励，例如「喝奶茶」「玩游戏 20 分钟」「买一本书」，并设置需要多少积分才能兑换。"),
    ])

    doc.add_heading("4. 额外功能：面板规则设定", level=1)
    para(doc, "当我希望面板按照自己的节奏运行时，可以进入面板规则。这里主要控制两件事：面板倒计时，以及隔天如何处理任务状态。")
    image(doc, "real_board_rules_marked.png", "面板规则：设置倒计时和任务重置方式")
    marker_table(doc, [
        ("1", "面板倒计时：我可以给整个面板设置一个完成时限。时间到后会出现超时提示，我可以选择把整个面板标记为完成、延期，或取消倒计时；它不是一到时间就固定清空任务。"),
        ("2", "隔天自动重置任务状态：任务文字会保留，但完成状态会重置为未完成。适合每天重复使用的习惯面板，例如晨间流程、睡前流程。"),
        ("3", "自定义重置时间：如果我不是以 24:00 作为一天结束，可以自定义每天重置的时间。"),
        ("4", "隔天自动清空格子任务：任务文字和完成状态都会清空。适合一次性清单、临时计划或只用一天的任务面板。"),
        ("5", "保持任务状态不自动重置：任务和完成状态都会保留，直到我手动修改。适合长期目标或不想每天刷新的面板。"),
    ])

    doc.add_heading("5. 额外功能：格子大小、多个面板和快速编辑", level=1)
    para(doc, "如果我有不同使用场景，可以添加多个面板。例如一个面板放早晨流程，一个面板放学习任务，另一个面板放睡前流程。切换面板后，每个场景都有自己的任务和进度。")
    para(doc, "如果我想一次整理很多任务，可以使用快速编辑。这里可以调整面板大小、管理任务分组、选择任务，并把预览结果应用到当前面板。")
    image(doc, "real_quick_edit_marked.png", "快速编辑：批量整理任务，并调整面板内容")
    marker_table(doc, [
        ("1", "格子大小调整：我可以调整面板尺寸。格子越少，任务越轻；格子越多，适合更完整的计划。"),
        ("2", "面板预览：我可以先看到任务会怎样排布，确认后再应用到首页。"),
        ("3", "任务分组：我可以把常用任务分组保存，例如「早晨」「学习」「家务」，之后可以重复使用。"),
        ("4", "任务列表：我可以勾选任务加入预览，也可以编辑或删除已有任务。"),
        ("5", "应用到面板：确认预览无误后，点击应用，当前面板就会更新成这套任务。"),
    ])

    doc.add_heading("6. 使用建议", level=1)
    bullets(doc, [
        "我可以为不同场景添加多个面板，例如早晨面板、学习面板、运动面板和睡前面板。",
        "任务越小，越容易开始。可以把「打扫房间」改成「整理桌面 5 分钟」。",
        "一个面板最好只放一种场景，这样打开时更清楚自己下一步要做什么。",
        "每日习惯建议选择「重置任务状态」；一次性清单才选择「清空格子任务」。",
        "自定义奖励最好具体、可兑换，这样积分才会真的有吸引力。",
    ])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
