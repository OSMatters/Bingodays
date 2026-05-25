from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "manual_assets"
OUT = ROOT / "Bingodays_产品使用手册.docx"


def font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size=size, index=0)
            except Exception:
                continue
    return ImageFont.load_default()


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "PingFang SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(9.5)
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_hyper_compact_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.35), Inches(4.95)]
    for label, detail in rows:
        cells = table.add_row().cells
        for i, width in enumerate(widths):
            cells[i].width = width
        set_cell_text(cells[0], label, True)
        set_cell_text(cells[1], detail)
        set_cell_shading(cells[0], "E8EEF5")
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.05)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.name = "PingFang SC"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        run.font.size = Pt(10.5)


def add_steps(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        run.font.name = "PingFang SC"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        run.font.size = Pt(10.5)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(85, 85, 85)
    run.font.name = "PingFang SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")


def add_image(doc: Document, path: Path, width: float, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def make_flow_image():
    out = ASSETS / "feature_flow.png"
    w, h = 1600, 720
    img = Image.new("RGB", (w, h), "#F7F8FA")
    draw = ImageDraw.Draw(img)
    title_font = font(48, True)
    box_font = font(32)
    small_font = font(24)
    draw.text((70, 50), "Bingodays 使用路径", fill="#17202C", font=title_font)

    boxes = [
        ("新建/选择面板", "管理不同生活场景"),
        ("填写任务", "长按格子或快速编辑"),
        ("完成任务", "点击格子点亮状态"),
        ("连成 Bingo", "获得积分与正反馈"),
        ("复盘与奖励", "日记、贴纸、自定义奖励"),
    ]
    x0, y0, bw, bh, gap = 70, 180, 250, 210, 55
    for i, (head, sub) in enumerate(boxes):
        x = x0 + i * (bw + gap)
        draw.rounded_rectangle((x, y0, x + bw, y0 + bh), radius=28, fill="#FFFFFF", outline="#D8DEE8", width=3)
        draw.rounded_rectangle((x, y0, x + bw, y0 + 18), radius=9, fill="#D9A770")
        draw.text((x + 28, y0 + 58), head, fill="#17202C", font=box_font)
        draw.text((x + 28, y0 + 118), sub, fill="#606A78", font=small_font)
        if i < len(boxes) - 1:
            ax = x + bw + 13
            ay = y0 + bh // 2
            draw.line((ax, ay, ax + gap - 26, ay), fill="#D9A770", width=8)
            draw.polygon([(ax + gap - 26, ay - 18), (ax + gap - 26, ay + 18), (ax + gap + 8, ay)], fill="#D9A770")

    note = "核心原则：把任务拆小、放进格子、用完成反馈推动下一步。"
    draw.rounded_rectangle((70, 500, 1530, 625), radius=24, fill="#EEF3F8", outline="#D8DEE8", width=2)
    draw.text((115, 542), note, fill="#17202C", font=font(34))
    img.save(out)
    return out


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 14, 8),
        ("Heading 2", 13, "2E74B5", 10, 5),
        ("Heading 3", 11.5, "1F4D78", 7, 3),
    ]:
        style = styles[name]
        style.font.name = "PingFang SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("Bingodays 产品使用手册")
    run.font.name = "PingFang SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 32, 44)

    p = doc.add_paragraph()
    run = p.add_run("适用于 BingoADHD 2 代码库 / iOS 应用 Bingodays 2.2.0")
    run.font.name = "PingFang SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(85, 85, 85)


def build():
    ASSETS.mkdir(exist_ok=True)
    flow = make_flow_image()
    doc = Document()
    configure_styles(doc)
    add_title(doc)

    doc.add_paragraph(
        "Bingodays 是一款把日常任务放进 Bingo 面板的行动辅助工具。用户通过点击格子完成任务、连成 Bingo、获得积分和奖励，适合把注意力分散、启动困难或容易拖延的任务拆成可见的小步骤。"
    )
    add_hyper_compact_table(
        doc,
        [
            ("核心对象", "面板、格子任务、快捷任务库、积分、奖励、日记、小组件。"),
            ("主要入口", "首页左上角菜单、右上角积分星星、面板下方分享/倒计时/编辑按钮。"),
            ("免费限制", "默认 1 个面板，最多 4×4；快捷任务最多 8 个、分组最多 3 个、每组最多 5 个；自定义奖励最多 3 个。"),
            ("Pro 权益", "多个面板、5×5 面板、不限快捷任务与分组、不限自定义奖励、完整历史任务、每日自定义重置时间。"),
        ],
    )
    add_image(doc, flow, 6.2, "图 1：Bingodays 的基本使用路径")

    doc.add_heading("1. 初次启动", level=1)
    doc.add_paragraph("首次打开应用会进入引导流程，之后会出现新手任务提示。引导的目标不是讲解所有功能，而是让用户先完成几个很小的任务，体验从点亮格子到形成 Bingo 的反馈。")
    add_steps(
        doc,
        [
            "完成 Onboarding 后进入首页。",
            "在欢迎提示里点击 OK，应用会放入“喝一杯水、站立 1 分钟、拉伸身体”等入门任务。",
            "点击任务格标记完成；完成若干任务后会出现阶段提示。",
            "完成第一条 Bingo 后，长按格子即可编辑自己的下一轮任务。",
        ],
    )
    add_image(doc, ASSETS / "01_home_or_onboarding.png", 2.35, "图 2：首次进入后的新手任务引导")

    doc.add_heading("2. 首页面板", level=1)
    doc.add_paragraph("首页是日常使用的主工作区。中间是 Bingo 面板，每个格子代表一个任务；上方可切换或新增面板；右上角显示当前可用积分。")
    add_image(doc, ASSETS / "04_home_seeded_fixed.png", 2.35, "图 3：已配置任务后的首页面板")
    add_hyper_compact_table(
        doc,
        [
            ("左上角菜单", "打开日记、2048 模式、联系我们、设置、小组件说明、主题颜色和 Pro 入口。"),
            ("积分星星", "进入积分与奖励页，可兑换贴纸、创建奖励、查看兑换历史。"),
            ("面板名称", "当前面板入口；长按面板区域可打开面板管理动作。"),
            ("+ 添加棋盘", "创建新的任务面板。免费版只能保留 1 个面板，更多面板需要 Pro。"),
            ("九宫格任务", "点击切换完成状态；长按单个格子可编辑、隐藏或删除任务。"),
            ("底部按钮", "左侧为模板分享；中间为面板倒计时；右侧进入编辑模式。"),
        ],
    )

    doc.add_heading("3. 添加和编辑任务", level=1)
    doc.add_paragraph("任务既可以在单个格子中编辑，也可以通过快速编辑批量维护。单格编辑适合临时任务，快速编辑适合常用清单和成组任务。")
    add_steps(
        doc,
        [
            "长按空白或已有任务格，打开任务编辑页。",
            "输入任务名称；任务文本会被限制在较短长度内，适合写成“喝水”“读 10 页”这类行动语句。",
            "可选择是否强制完成。强制任务会让其他任务在它完成前处于受限状态，用于必须先做的任务。",
            "设置常驻日期：始终显示、自定义星期、仅当天。",
            "按需设置开始时间、结束时间或任务倒计时。",
            "保存后返回首页，任务会显示在对应格子里。",
        ],
    )
    add_bullets(
        doc,
        [
            "开始时间未到时，任务会显示为未激活状态；结束时间到了会进入超时处理。",
            "任务倒计时结束后，可以选择标记完成、删除、延期或替换任务。",
            "删除任务会出现确认，避免误操作。",
        ],
    )

    doc.add_heading("4. 完成任务和 Bingo 反馈", level=1)
    add_bullets(
        doc,
        [
            "点击任务格即可在完成与未完成之间切换。",
            "完成后的格子会变为主题色，并显示勾选反馈。",
            "连成一行、一列或对角线会记录 Bingo 进度，并触发积分结算与完成动画。",
            "积分结算以完成格子数、完成线数和全盘完成奖励为基础。",
            "当天任务会同步写入 Bingo 日记，便于后续复盘。",
        ],
    )

    doc.add_heading("5. 快速编辑 My Tasks", level=1)
    doc.add_paragraph("快速编辑页用于维护可复用任务库，并把任务批量应用到当前面板。它支持普通任务、任务分组、历史任务复用、预览格子、拖拽调整和应用分组。")
    add_hyper_compact_table(
        doc,
        [
            ("普通任务", "保存单个常用任务，之后可快速加入面板。"),
            ("任务分组", "把一组例行事项保存为模板，例如“晨间启动”“睡前收尾”。"),
            ("应用分组", "将某个分组按预览位置批量填入面板。"),
            ("历史任务", "从过去使用过的任务中复用；免费版只显示部分历史任务。"),
            ("开始日期", "任务可设置月/日开始显示，适合季节性或预约类任务。"),
        ],
    )

    doc.add_heading("6. 面板管理和规则", level=1)
    add_bullets(
        doc,
        [
            "面板可重命名、删除、切换，并支持左右滑动切换。",
            "编辑模式中可以调整格子数量、随机排序、清空面板、进入快速编辑。",
            "免费版最大 4×4，Pro 可解锁 5×5。",
            "面板规则可设置次日自动重置完成状态、次日清空所有任务、保持任务状态不变，或在 Pro 中自定义每日重置时间。",
            "面板倒计时可按天、小时、分钟设置；倒计时结束后可清空或延期。",
        ],
    )

    doc.add_heading("7. 积分、贴纸和奖励", level=1)
    add_bullets(
        doc,
        [
            "右上角星星显示当前可用积分。",
            "积分页包含贴纸和我的奖励两个主要标签。",
            "贴纸可用积分兑换，已拥有的贴纸可添加到首页作为装饰。",
            "自定义奖励可设置标题、所需积分、emoji 和主题图；兑换后会写入兑换历史。",
            "免费版最多创建 3 个自定义奖励，Pro 不限数量。",
        ],
    )

    doc.add_heading("8. Bingo 日记", level=1)
    doc.add_paragraph("Bingo 日记用于查看完成记录和超时记录。它会汇总任务完成次数、近期活动、周/月/年统计，并可查看某一天的面板快照。")
    add_hyper_compact_table(
        doc,
        [
            ("日历视图", "点击已完成日期，查看当天的 Bingo 面板。"),
            ("任务统计", "查看某个任务累计完成次数、本周完成次数和近期记录。"),
            ("超时任务", "单独统计被倒计时或时间窗口判定为超时的任务。"),
            ("连续天数", "侧边栏展示 streak 和阶段目标，用于长期反馈。"),
        ],
    )

    doc.add_heading("9. 2048 模式", level=1)
    doc.add_paragraph("2048 模式是带主题的任务游戏模式。用户选择主题后，完成任务格，再按 2048 规则滑动合并已完成格子，把任务进度转化为游戏反馈。")
    add_bullets(
        doc,
        [
            "进入入口：左上角菜单 → 2048。",
            "点击格子标记任务完成。",
            "滑动棋盘移动并合并已完成格子。",
            "刷新未完成任务时，只刷新未完成且未合并的格子。",
            "退出或切换主题前会出现确认，避免误丢进度。",
        ],
    )

    doc.add_heading("10. 小组件和模板分享", level=1)
    add_bullets(
        doc,
        [
            "小组件包含面板小组件和倒计时小组件，可在系统桌面快速查看当前面板或倒计时。",
            "模板分享会把当前面板任务编码为分享链接或图片，便于把一套任务模板发给他人。",
            "导入模板时，Pro 用户可作为新面板导入；免费用户通常会替换当前面板。",
            "空面板不能生成有效模板。",
        ],
    )

    doc.add_heading("11. 设置和会员", level=1)
    add_hyper_compact_table(
        doc,
        [
            ("触感反馈", "控制点击、完成、强调反馈的震动效果。"),
            ("音效", "控制完成、得分等音效。"),
            ("主题颜色", "切换首页和组件的主题色。"),
            ("联系我们", "打开邮件反馈入口。"),
            ("Pro/会员中心", "查看订阅、终身会员、恢复购买和权益说明。"),
        ],
    )

    doc.add_heading("12. 常见问题", level=1)
    add_hyper_compact_table(
        doc,
        [
            ("为什么只能创建一个面板？", "免费版限制为 1 个面板；开通 Pro 后可创建多个面板。"),
            ("为什么不能用 5×5？", "免费版最多 4×4；5×5 是 Pro 功能。"),
            ("任务没有显示怎么办？", "检查是否设置了开始日期、结束时间、仅当天或星期重复规则。"),
            ("倒计时结束后怎么办？", "按弹窗选择完成、延期、删除或替换任务。"),
            ("积分不够兑换奖励？", "继续完成任务和 Bingo 线；兑换会扣除可用积分，但历史积分仍用于长期统计。"),
            ("如何恢复购买？", "进入 Pro/会员中心，使用恢复订阅入口同步 App Store 购买状态。"),
        ],
    )

    doc.add_heading("附录：代码依据", level=1)
    doc.add_paragraph(
        "本手册根据 BingoADHD 2 当前代码整理，重点参考 ContentView、BingoViewModel、BingoModels、Localization、QuickEdit、BoardRule、PointsFlow、BingoDiary、BlackBoxMode、Widget 相关文件。截图来自本地 iPhone 16 iOS 18.5 模拟器构建运行结果。"
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
