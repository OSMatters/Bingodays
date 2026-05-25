from __future__ import annotations

from pathlib import Path
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "manual_assets"
OUT = ROOT / "Bingodays_产品使用手册_案例版.docx"

INK = "#17202C"
MUTED = "#667085"
ACCENT = "#D9A770"
BLUE = "#2E74B5"
BG = "#F6F7F9"
CARD = "#FFFFFF"
BORDER = "#D8DEE8"


def pil_font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size=size, index=0)
            except Exception:
                continue
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int):
    lines: list[str] = []
    for para in text.split("\n"):
        current = ""
        for char in para:
            trial = current + char
            if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = char
        if current:
            lines.append(current)
    return lines


def draw_textbox(draw, xy, text, font, fill, max_width, line_gap=8):
    x, y = xy
    for line in wrap_text(draw, text, font, max_width):
        draw.text((x, y), line, font=font, fill=fill)
        y += font.size + line_gap
    return y


def marker(draw, x, y, n: int):
    r = 24
    draw.ellipse((x - r, y - r, x + r, y + r), fill=BLUE)
    txt = str(n)
    bb = draw.textbbox((0, 0), txt, font=pil_font(25, True))
    draw.text((x - (bb[2] - bb[0]) / 2, y - (bb[3] - bb[1]) / 2 - 2), txt, fill="white", font=pil_font(25, True))


def callout(draw, x, y, n, title, body, width=390):
    marker(draw, x, y, n)
    draw_textbox(draw, (x + 42, y - 26), title, pil_font(26, True), INK, width)
    draw_textbox(draw, (x + 42, y + 12), body, pil_font(21), MUTED, width)


def rounded(draw, box, radius=24, fill=CARD, outline=BORDER, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def base_canvas(w=1600, h=1100, title=""):
    img = Image.new("RGB", (w, h), BG)
    draw = ImageDraw.Draw(img)
    if title:
        draw.text((70, 48), title, font=pil_font(48, True), fill=INK)
    return img, draw


def phone_frame(draw, x, y, w=520, h=930):
    rounded(draw, (x, y, x + w, y + h), 54, "#FDFDFD", "#EAECF0", 3)
    draw.rounded_rectangle((x + 170, y + 18, x + 350, y + 54), radius=20, fill="#111111")


def make_home_case_image():
    img, draw = base_canvas(title="用户案例：从任务编辑到 Bingo 再到奖励")
    phone_frame(draw, 80, 125, 500, 850)
    px, py = 120, 230
    draw.text((px, py), "晨间启动", font=pil_font(36, True), fill=INK)
    draw.text((px + 340, py), "星 12", font=pil_font(28, True), fill=ACCENT)
    grid_x, grid_y, cell, gap = px, py + 90, 130, 18
    tasks = [("喝水", True), ("站立", True), ("拉伸", True), ("计划", False), ("整理", False), ("番茄钟", False), ("阅读", False), ("散步", False), ("睡觉", False)]
    for i, (text, done) in enumerate(tasks):
        r, c = divmod(i, 3)
        x = grid_x + c * (cell + gap)
        y = grid_y + r * (cell + gap)
        fill = ACCENT if done else "#FFFFFF"
        rounded(draw, (x, y, x + cell, y + cell), 20, fill, "#E4E7EC", 2)
        draw_textbox(draw, (x + 18, y + 42), text, pil_font(29, True), "white" if done else INK, cell - 36)
    draw.line((grid_x, grid_y + cell / 2, grid_x + 3 * cell + 2 * gap, grid_y + cell / 2), fill="#FFFFFF", width=7)
    draw.text((px, py + 625), "分享   倒计时   编辑", font=pil_font(24, True), fill=ACCENT)

    callout(draw, 670, 190, 1, "编辑任务", "长按任务格打开编辑弹窗；先写小任务，再设置重复、时间窗口或倒计时。")
    callout(draw, 670, 380, 2, "点击完成", "完成一个任务就点亮一个格子，形成行、列或对角线后触发 Bingo。")
    callout(draw, 670, 570, 3, "获得积分", "积分来自完成格子、完成线和全盘完成奖励；右上角星星显示可用积分。")
    callout(draw, 670, 760, 4, "兑换奖励", "进入积分页，兑换贴纸或创建自定义奖励，例如“30 分钟游戏时间”。")
    img.save(ASSETS / "case_home_flow.png")


def make_rules_image():
    img, draw = base_canvas(title="面板规则：决定任务如何跨天保留或刷新")
    phone_frame(draw, 70, 125, 560, 900)
    x, y = 115, 220
    draw.text((x, y), "面板规则", font=pil_font(40, True), fill=INK)
    y += 80
    rounded(draw, (x, y, x + 470, y + 142), 22, "#FFFFFF")
    draw.text((x + 26, y + 22), "面板倒计时", font=pil_font(28, True), fill=INK)
    draw.text((x + 26, y + 70), "开启后：到点清空面板任务", font=pil_font(23), fill=MUTED)
    marker(draw, x + 430, y + 50, 1)
    y += 190
    options = [
        ("隔天自动重置任务状态", "保留任务文字，只把完成状态变回未完成", 2),
        ("自定义重置任务状态时间", "Pro；例如每天 03:00 重置，适合夜班或晚睡用户", 3),
        ("隔天自动清空格子任务", "第二天清空所有格子，适合一次性计划", 4),
        ("保持任务状态不自动重置", "任务状态一直保留，直到用户手动修改", 5),
    ]
    for title, body, n in options:
        rounded(draw, (x, y, x + 470, y + 118), 20, "#FFFFFF")
        draw.ellipse((x + 26, y + 33, x + 56, y + 63), outline=ACCENT, width=4)
        draw.text((x + 74, y + 22), title, font=pil_font(24, True), fill=INK)
        draw_textbox(draw, (x + 74, y + 58), body, pil_font(19), MUTED, 330)
        marker(draw, x + 430, y + 48, n)
        y += 135

    callout(draw, 710, 185, 1, "倒计时是一次性清空", "它不等同于每日规则。适合“本周冲刺”“今天必须完成”的面板。")
    callout(draw, 710, 360, 2, "默认适合日常习惯", "任务文字不丢，第二天只重置完成状态，适合喝水、运动、阅读。")
    callout(draw, 710, 535, 3, "自定义重置跟着作息走", "只有选择“重置任务状态”时才出现；未开 Pro 时保存会进入会员页。")
    callout(draw, 710, 710, 4, "清空适合临时清单", "例如搬家、旅行准备、会议当天待办，第二天不需要保留。")
    callout(draw, 710, 885, 5, "不重置适合项目进度", "例如论文、证书学习、长期清单，完成状态需要持续保留。")
    img.save(ASSETS / "rules_annotated.png")


def make_task_edit_image():
    img, draw = base_canvas(title="长按任务格编辑：单个任务的行为逻辑")
    phone_frame(draw, 70, 125, 560, 900)
    x, y = 115, 200
    draw.text((x, y), "编辑任务", font=pil_font(38, True), fill=INK)
    y += 70
    fields = [
        ("任务名称 + 语音输入", "最多约 20 字；适合写成可执行动作", 1),
        ("开始时间 / 结束时间", "控制任务何时显示、何时超时", 2),
        ("强制完成", "开启后，其他任务会被锁定直到它完成", 3),
        ("常驻日期", "始终 / 自定义星期 / 仅当天", 4),
        ("任务倒计时", "15/30/45/60 分钟或自定义分钟", 5),
        ("保存 / 删除 / 关闭", "校验通过后保存；删除需确认", 6),
    ]
    for title, body, n in fields:
        rounded(draw, (x, y, x + 470, y + 102), 18, "#FFFFFF")
        draw.text((x + 22, y + 17), title, font=pil_font(24, True), fill=INK)
        draw.text((x + 22, y + 54), body, font=pil_font(19), fill=MUTED)
        marker(draw, x + 430, y + 42, n)
        y += 118

    callout(draw, 710, 185, 1, "先写动作，不写愿望", "把“保持健康”改成“散步 10 分钟”，完成判断更明确。")
    callout(draw, 710, 335, 2, "时间窗口决定显示和超时", "开始时间未到时不会进入普通待办；结束时间到后进入超时处理。")
    callout(draw, 710, 485, 3, "强制完成用于关键前置任务", "例如“吃药”“提交日报”；它会让其他格子暂时不可点。")
    callout(draw, 710, 635, 4, "重复规则决定哪天出现", "“仅当天”适合一次性任务；星期选择适合固定例行。")
    callout(draw, 710, 785, 5, "倒计时用于启动任务", "适合番茄钟、限时收拾、专注工作。已完成任务不能再新增倒计时。")
    callout(draw, 710, 935, 6, "保存有校验", "结束必须晚于开始；开始/结束不能早于当前时间。")
    img.save(ASSETS / "task_edit_annotated.png")


def make_quick_edit_image():
    img, draw = base_canvas(title="快速编辑：把任务库批量应用到面板")
    phone_frame(draw, 70, 125, 620, 900)
    x, y = 115, 200
    draw.text((x, y), "快速编辑", font=pil_font(38, True), fill=INK)
    y += 70
    rounded(draw, (x, y, x + 530, y + 72), 18, "#FFFFFF")
    draw.text((x + 24, y + 20), "全部任务  |  晨间启动  |  + 分组", font=pil_font(24, True), fill=INK)
    marker(draw, x + 500, y + 34, 1)
    y += 105
    rounded(draw, (x, y, x + 530, y + 210), 22, "#FFFFFF")
    draw.text((x + 24, y + 22), "任务库", font=pil_font(27, True), fill=INK)
    for i, t in enumerate(["喝水", "站立 1 分钟", "拉伸身体"]):
        yy = y + 66 + i * 43
        draw.ellipse((x + 28, yy + 5, x + 50, yy + 27), outline=ACCENT, width=3)
        draw.text((x + 66, yy), t, font=pil_font(22), fill=INK)
    marker(draw, x + 500, y + 85, 2)
    y += 245
    rounded(draw, (x, y, x + 530, y + 210), 22, "#FFFFFF")
    draw.text((x + 24, y + 22), "面板预览", font=pil_font(27, True), fill=INK)
    cell, gap = 72, 12
    for i in range(9):
        r, c = divmod(i, 3)
        xx = x + 28 + c * (cell + gap)
        yy = y + 68 + r * (cell + gap)
        rounded(draw, (xx, yy, xx + cell, yy + cell), 10, ACCENT if i < 3 else "#F2F4F7", "#E4E7EC", 1)
    draw.text((x + 315, y + 80), "3 / 9", font=pil_font(28, True), fill=ACCENT)
    marker(draw, x + 500, y + 120, 3)
    y += 245
    rounded(draw, (x, y, x + 255, y + 70), 16, ACCENT, ACCENT)
    draw.text((x + 70, y + 20), "应用", font=pil_font(25, True), fill="white")
    rounded(draw, (x + 275, y, x + 530, y + 70), 16, "#FFFFFF", ACCENT, 2)
    draw.text((x + 350, y + 20), "取消", font=pil_font(25, True), fill=INK)
    marker(draw, x + 500, y + 32, 4)

    callout(draw, 760, 185, 1, "筛选和分组", "全部任务包含历史任务；分组用于保存一套可复用清单。免费版最多 3 个分组。")
    callout(draw, 760, 365, 2, "选择任务进入预览", "勾选任务后会放入预览格；未到开始日期的任务不会被加入。")
    callout(draw, 760, 545, 3, "预览就是即将应用的面板", "可调整目标格子数，拖拽预览格改变位置，最多填满当前网格。")
    callout(draw, 760, 725, 4, "应用才会写回首页", "退出前若有未保存改动，会提示是否放弃。")
    callout(draw, 760, 905, 5, "免费版容量限制", "普通快捷任务最多 8 个；每组最多 5 个任务；Pro 取消限制并可用 5×5。")
    img.save(ASSETS / "quick_edit_annotated.png")


def make_reward_image():
    img, draw = base_canvas(title="积分与自定义奖励：把完成感换成可见回报", h=900)
    rounded(draw, (80, 160, 1520, 760), 34, "#FFFFFF")
    cols = [
        ("1 完成任务", "点击格子完成任务，连成 Bingo 后获得更多积分。"),
        ("2 进入积分页", "点击首页右上角星星，查看可用积分。"),
        ("3 创建奖励", "设置奖励名称、所需积分、emoji 或主题图。"),
        ("4 兑换奖励", "积分足够时兑换，记录会进入兑换历史。"),
    ]
    for i, (title, body) in enumerate(cols):
        x = 130 + i * 350
        rounded(draw, (x, 250, x + 280, 560), 24, "#F7F8FA", "#E4E7EC")
        draw.text((x + 28, 300), title, font=pil_font(30, True), fill=INK)
        draw_textbox(draw, (x + 28, 365), body, pil_font(23), MUTED, 220)
        if i < 3:
            draw.line((x + 292, 405, x + 330, 405), fill=ACCENT, width=8)
            draw.polygon([(x + 330, 385), (x + 330, 425), (x + 365, 405)], fill=ACCENT)
    draw.text((130, 640), "案例：小林设置“30 分钟游戏时间 = 30 积分”。当晨间面板完成一条 Bingo 后，进入积分页兑换奖励，形成正反馈。", font=pil_font(28), fill=INK)
    img.save(ASSETS / "reward_flow.png")


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_docx_run_font(run, size=10.5, bold=False, color=None):
    run.font.name = "PingFang SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for label, detail in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.45)
        cells[1].width = Inches(4.85)
        set_cell_shading(cells[0], "E8EEF5")
        for cell, text, bold in [(cells[0], label, True), (cells[1], detail, False)]:
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_docx_run_font(run, 9.5, bold)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_docx_run_font(run)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_docx_run_font(run)


def add_para(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_docx_run_font(run)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_docx_run_font(r, 9, False, "555555")
    r.italic = True


def add_image(doc, filename, width=6.3, caption=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    if caption:
        add_caption(doc, caption)


def configure_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 11.5, "1F4D78")]:
        style = doc.styles[name]
        style.font.name = "PingFang SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Bingodays 产品使用手册：案例与功能逻辑版")
    set_docx_run_font(r, 23, True, "17202C")
    p = doc.add_paragraph()
    r = p.add_run("从产品定位、用户案例、面板规则、任务编辑到快速编辑的完整说明")
    set_docx_run_font(r, 10, False, "555555")


def build():
    ASSETS.mkdir(exist_ok=True)
    make_home_case_image()
    make_rules_image()
    make_task_edit_image()
    make_quick_edit_image()
    make_reward_image()

    doc = Document()
    configure_doc(doc)
    title(doc)

    doc.add_heading("1. 产品定位与适用人群", level=1)
    add_para(doc, "Bingodays 的定位是“把任务管理做成可完成、可反馈、可奖励的 Bingo 面板”。它不是传统待办清单，也不是复杂项目管理工具，而是面向日常行动启动、习惯坚持和轻量复盘的个人任务产品。")
    add_table(doc, [
        ("产品核心", "把任务拆成格子，通过点击完成、连线 Bingo、积分奖励和日记复盘形成闭环。"),
        ("适用人群", "启动困难、注意力容易分散、容易被长清单压垮、需要即时反馈的人；也适合希望建立运动、学习、家务、睡前流程等习惯的人。"),
        ("不适合", "多人协作、严格项目排期、复杂依赖管理、企业级任务审批。"),
        ("设计原则", "任务要小、反馈要快、奖励要具体、规则要跟着用户作息走。"),
    ])

    doc.add_heading("2. 用户案例：小林如何从编辑到 Bingo，再用积分换奖励", level=1)
    add_para(doc, "小林是一名产品运营。她早上经常知道自己该做什么，但很难启动。她用 Bingodays 建了一个“晨间启动”面板，把早晨要做的事情拆成 9 个格子。")
    add_image(doc, "case_home_flow.png", 6.35, "图 1：用户案例的完整使用路径")
    add_steps(doc, [
        "新建或选择“晨间启动”面板。",
        "长按空白格，依次编辑“喝水、站立 1 分钟、拉伸身体、写今日计划”等任务。",
        "完成任务后点击格子。前三个任务在同一行，全部完成后形成一条 Bingo。",
        "进入右上角积分页，看到完成任务和 Bingo 带来的积分变化。",
        "在“我的奖励”中创建自定义奖励：名称“30 分钟游戏时间”，所需积分 30。",
        "积分足够后兑换奖励，记录进入兑换历史；之后小林更愿意继续完成下一轮任务。",
    ])
    add_image(doc, "reward_flow.png", 6.35, "图 2：积分和自定义奖励的使用闭环")

    doc.add_heading("3. 面板规则设定逻辑", level=1)
    add_para(doc, "面板规则解决的是一个核心问题：任务跨天之后，到底应该保留什么、重置什么、清空什么。不同用户场景应该选择不同规则。")
    add_image(doc, "rules_annotated.png", 6.35, "图 3：面板规则页面功能标注")
    add_table(doc, [
        ("1 面板倒计时", "独立于每日重置规则。开启后设置天/小时/分钟，到点后面板任务会清空，适合限时冲刺。"),
        ("2 自动重置状态", "默认习惯场景最适合。第二天保留任务文字，只把已完成状态变回未完成。"),
        ("3 自定义重置时间", "只在“自动重置状态”模式下生效，属于 Pro 功能。适合夜班、晚睡或非零点作息。"),
        ("4 自动清空任务", "第二天清空所有格子任务，适合一次性任务面板，例如旅行打包、会议准备。"),
        ("5 保持不重置", "完成状态和任务都保留，适合长期项目或阶段性学习计划。"),
    ])
    doc.add_heading("小林如何选择", level=2)
    add_bullets(doc, [
        "晨间习惯面板：选择“隔天自动重置任务状态”，每天任务相同，只重置完成状态。",
        "周末大扫除面板：选择“隔天自动清空格子任务”，任务做完后第二天不再需要。",
        "论文推进面板：选择“保持任务状态不自动重置”，让长期进度一直保留。",
        "晚睡作息面板：如果开通 Pro，把自定义重置时间设为 03:00，避免 00:00 把当天任务提前重置。",
    ])

    doc.add_heading("4. 长按任务格子编辑弹窗逻辑", level=1)
    add_para(doc, "长按任务格子用于编辑单个任务。它适合精细控制一个任务什么时候出现、是否必须先完成、是否重复、是否需要倒计时。")
    add_image(doc, "task_edit_annotated.png", 6.35, "图 4：长按任务格编辑弹窗功能标注")
    add_table(doc, [
        ("1 任务名称", "任务文本会限制长度。建议写成动作，而不是目标，例如“读 10 页”优于“学习”。麦克风用于语音输入。"),
        ("2 开始/结束时间", "开始时间控制任务何时显示；结束时间到达后进入超时处理。结束必须晚于开始，且不能早于当前时间。"),
        ("3 强制完成", "开启后其他任务被锁定，直到该任务完成。适合吃药、提交日报、打卡等前置任务。"),
        ("4 常驻日期", "始终显示、按星期重复、仅当天。仅当天适合一次性任务；星期重复适合固定例行。"),
        ("5 任务倒计时", "用于任务开始后的限时专注。可选 15/30/45/60 分钟或自定义分钟。"),
        ("6 保存/删除/关闭", "保存前会校验时间和倒计时逻辑；删除任务需要确认；关闭不会保存未提交变更。"),
    ])
    doc.add_heading("小林的设置示例", level=2)
    add_bullets(doc, [
        "“写今日计划”：始终显示，不设倒计时，作为每天固定启动任务。",
        "“番茄钟 25 分钟”：设置任务倒计时 25 分钟，到点后提醒处理。",
        "“吃早餐”：设为强制完成，避免她先刷手机再忘记早餐。",
        "“周三倒垃圾”：选择自定义星期，只在周三出现。",
    ])

    doc.add_heading("5. 快速编辑页面功能逻辑", level=1)
    add_para(doc, "快速编辑页面用于批量管理任务库，并把一组任务应用到当前面板。它适合“我有一套常用流程，想快速放到面板里”的场景。")
    add_image(doc, "quick_edit_annotated.png", 6.35, "图 5：快速编辑页面功能标注")
    add_table(doc, [
        ("1 筛选与分组", "全部任务显示历史和任务库；分组保存成套任务。免费版最多 3 个分组，Pro 不限。"),
        ("2 任务库行", "勾选代表加入预览；可编辑任务名、删除任务，也可新增任务。"),
        ("3 面板预览", "预览就是即将应用到首页的格子布局。任务可拖拽换位置，超出当前网格容量的任务不会继续加入。"),
        ("4 应用/取消", "点击应用才会写回首页；取消或退出时如有未保存改动，会出现放弃确认。"),
        ("5 容量限制", "免费版快捷任务最多 8 个、分组最多 3 个、每组最多 5 个；Pro 可解锁更多容量和 5×5。"),
    ])
    doc.add_heading("小林如何用快速编辑", level=2)
    add_steps(doc, [
        "创建分组“晨间启动”。",
        "在分组里添加“喝水、站立 1 分钟、拉伸身体、写今日计划”。",
        "勾选这些任务，让它们进入右侧或上方的面板预览。",
        "拖拽预览格，把最容易启动的任务放在第一行。",
        "点击应用，当前面板立即替换为这套晨间任务。",
    ])

    doc.add_heading("结论：如何把产品用好", level=1)
    add_bullets(doc, [
        "先从 3 个很小的任务开始，不要一开始填满 5×5。",
        "习惯类面板用“自动重置状态”，项目类面板用“保持不重置”，一次性清单用“自动清空”。",
        "长按编辑解决单个任务的精细规则，快速编辑解决成组任务的批量复用。",
        "自定义奖励要具体、可兑换、低门槛，例如“咖啡 1 杯”“休息 20 分钟”。",
        "每天只追求形成一条 Bingo，而不是追求完美全盘完成。",
    ])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
