from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "manual_assets"
OUT = ROOT / "Bingodays_产品使用手册_真实截图版.docx"


def pil_font(size: int):
    for path in [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size=size, index=0)
            except Exception:
                pass
    return ImageFont.load_default()


def marker(draw: ImageDraw.ImageDraw, x: int, y: int, n: int):
    r = 30
    draw.ellipse((x - r, y - r, x + r, y + r), fill="#2E74B5", outline="white", width=5)
    text = str(n)
    font = pil_font(30)
    bb = draw.textbbox((0, 0), text, font=font)
    draw.text((x - (bb[2] - bb[0]) / 2, y - (bb[3] - bb[1]) / 2 - 2), text, fill="white", font=font)


def annotate_image(source: str, target: str, points: list[tuple[int, int, int]]):
    img = Image.open(ASSETS / source).convert("RGB")
    draw = ImageDraw.Draw(img)
    for x, y, n in points:
        marker(draw, x, y, n)
    img.save(ASSETS / target)


def make_annotated_images():
    annotate_image(
        "04_home_seeded_fixed.png",
        "real_home_marked.png",
        [(105, 260, 1), (790, 260, 2), (150, 590, 3), (795, 590, 4), (470, 1040, 5), (845, 1690, 6)],
    )
    annotate_image(
        "real_board_rules.png",
        "real_board_rules_marked.png",
        [(790, 955, 1), (83, 1173, 2), (147, 1338, 3), (82, 1450, 4), (82, 1562, 5)],
    )
    annotate_image(
        "real_task_edit.png",
        "real_task_edit_marked.png",
        [(800, 336, 1), (120, 520, 2), (160, 892, 3), (125, 1128, 4), (120, 1585, 5), (500, 1980, 6)],
    )
    annotate_image(
        "real_quick_edit.png",
        "real_quick_edit_marked.png",
        [(760, 380, 1), (475, 800, 2), (145, 1370, 3), (185, 1650, 4), (850, 188, 5)],
    )
    annotate_image(
        "real_points.png",
        "real_points_marked.png",
        [(500, 350, 1), (240, 570, 2), (850, 220, 3), (220, 875, 4), (690, 570, 5)],
    )


def set_run(run, size=10.5, bold=False, color=None):
    run.font.name = "PingFang SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = doc.styles[name]
        style.font.name = "PingFang SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("2E74B5")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def para(doc, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run(r)


def steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run(r)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table(doc, rows):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    tbl.autofit = False
    for label, detail in rows:
        cells = tbl.add_row().cells
        cells[0].width = Inches(1.35)
        cells[1].width = Inches(4.95)
        shade(cells[0], "E8EEF5")
        for cell, text, bold in [(cells[0], label, True), (cells[1], detail, False)]:
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run(r, 9.5, bold)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph()


def image(doc, file, caption, width=2.6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / file), width=Inches(width))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    set_run(r, 9, False, "555555")
    r.italic = True


def title(doc):
    p = doc.add_paragraph()
    r = p.add_run("Bingodays 产品使用手册：真实截图标注版")
    set_run(r, 23, True, "17202C")
    p = doc.add_paragraph()
    r = p.add_run("使用产品内真实界面截图，并在截图上标记功能点")
    set_run(r, 10, False, "555555")


def build():
    make_annotated_images()
    doc = Document()
    configure(doc)
    title(doc)

    doc.add_heading("1. 产品定位与适用人群", level=1)
    para(doc, "Bingodays 是一款把日常任务做成 Bingo 面板的行动辅助产品。它把任务拆成格子，让用户通过点击完成、连线 Bingo、获得积分和兑换奖励，形成“行动-反馈-奖励-复盘”的闭环。")
    table(doc, [
        ("适用人群", "启动困难、注意力容易分散、容易被长清单压垮、需要即时反馈的人。"),
        ("典型场景", "晨间启动、睡前流程、家务清单、学习计划、运动习惯、短期冲刺。"),
        ("不适合", "多人协作、复杂项目排期、严格审批或跨团队任务管理。"),
    ])

    doc.add_heading("2. 用户案例：从编辑任务到 Bingo，再到积分奖励", level=1)
    para(doc, "小林早上很难启动，于是建立“晨间启动”面板，把早晨流程拆成 9 个小任务。")
    image(doc, "real_home_marked.png", "图 1：真实首页截图与基础功能标注")
    table(doc, [
        ("1 菜单", "进入日记、2048、设置、小组件说明和会员入口。"),
        ("2 积分入口", "点击星星进入积分页，查看贴纸、自定义奖励和兑换历史。"),
        ("3 当前面板", "显示当前面板名称；可切换、管理或新增面板。"),
        ("4 添加面板", "创建新的 Bingo 面板；免费版只能保留一个面板。"),
        ("5 任务格", "点击切换完成状态；连成行、列或对角线即形成 Bingo。"),
        ("6 编辑入口", "进入面板编辑能力，包括快速编辑、格子大小、随机排序和清空。"),
    ])
    steps(doc, [
        "长按空白格或已有任务格，编辑“喝水、站立 1 分钟、拉伸身体”等小任务。",
        "完成任务后点击格子，格子点亮并计入当天进度。",
        "同一行三个任务完成后形成 Bingo，获得正反馈和积分。",
        "点击右上角星星进入积分页，兑换贴纸或创建自定义奖励。",
    ])
    image(doc, "real_points_marked.png", "图 2：真实积分页截图与奖励功能标注")
    table(doc, [
        ("1 我的积分", "显示当前可用积分。"),
        ("2 贴纸", "用积分兑换贴纸，已拥有贴纸可添加到首页。"),
        ("3 历史", "查看积分兑换记录。"),
        ("4 贴纸卡片", "显示每个贴纸的所需积分和拥有状态。"),
        ("5 自定义奖励", "创建奖励名称、所需积分、emoji/主题图，积分足够时兑换。"),
    ])

    doc.add_heading("3. 面板规则设定逻辑", level=1)
    para(doc, "面板规则决定跨天后任务如何处理：保留文字、重置状态、清空任务，或完全不自动处理。")
    image(doc, "real_board_rules_marked.png", "图 3：真实面板规则截图与功能标注")
    table(doc, [
        ("1 面板倒计时", "一次性倒计时。开启后设置天/小时/分钟，到点后清空面板任务。"),
        ("2 自动重置状态", "保留任务文字，只把完成状态在 24:00 重置为未完成。适合每日习惯。"),
        ("3 自定义重置时间", "Pro 功能；仅在自动重置状态下使用，适合非零点作息。"),
        ("4 自动清空任务", "隔天清空所有格子任务。适合一次性清单。"),
        ("5 保持不重置", "任务和完成状态都保留，直到用户手动修改。适合长期项目。"),
    ])

    doc.add_heading("4. 长按任务格编辑弹窗逻辑", level=1)
    para(doc, "长按任务格打开单任务编辑弹窗。它负责定义一个任务的名称、显示时间、重复规则、是否强制完成和倒计时。")
    image(doc, "real_task_edit_marked.png", "图 4：真实任务编辑弹窗截图与功能标注")
    table(doc, [
        ("1 名称/语音", "输入任务名称，或用麦克风语音输入。建议写成可执行动作。"),
        ("2 时间窗口", "开始时间控制何时显示；结束时间到达后进入超时处理。"),
        ("3 强制完成", "开启后其他任务会锁定，直到该任务完成。"),
        ("4 重复天数", "始终、自定义星期或仅当天。"),
        ("5 倒计时", "设置任务倒计时，可选预设分钟或自定义分钟。"),
        ("6 保存", "保存前校验任务名、开始/结束时间和倒计时。"),
    ])

    doc.add_heading("5. 快速编辑页面功能逻辑", level=1)
    para(doc, "快速编辑用于把任务库批量应用到当前面板，适合维护一套可复用流程。")
    image(doc, "real_quick_edit_marked.png", "图 5：真实快速编辑截图与功能标注")
    table(doc, [
        ("1 面板大小", "调整目标面板尺寸；免费版最大 4×4，Pro 支持 5×5。"),
        ("2 面板预览", "预览即将应用到首页的任务布局。"),
        ("3 分组入口", "全部任务、任务分组和新增分组。免费版最多 3 个分组。"),
        ("4 任务列表", "勾选加入预览；可删除或编辑历史任务。"),
        ("5 应用", "点击后把预览任务写回首页；退出前有未保存改动会提示。"),
    ])

    doc.add_heading("结论", level=1)
    bullets(doc, [
        "日常习惯优先使用“自动重置任务状态”。",
        "一次性计划使用“自动清空格子任务”。",
        "长期项目使用“保持任务状态不自动重置”。",
        "单个任务的细节用长按编辑解决；成组任务复用用快速编辑解决。",
        "奖励设置要低门槛、具体、可兑换，才能形成持续反馈。",
    ])
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
